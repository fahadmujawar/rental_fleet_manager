"""Import the TCS machine register from Excel into the machines table.

Default is a dry run: the script reads the database and the Excel file,
prints a summary and writes data/import_report.csv, but writes nothing to
the database. With --apply it writes everything in one transaction.
"""

import argparse
import csv
import datetime
import os
import re
import sys
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import psycopg
from dotenv import load_dotenv

# 1. File and company
SCRIPT_DIR = Path(__file__).resolve().parent
DATA_DIR = SCRIPT_DIR.parent.parent / "data"
REGISTER_FILE = DATA_DIR / "Total_Assets.xlsx"
REGISTER_SHEET = "Sheet1"
REPORT_FILE = DATA_DIR / "import_report.csv"
COMPANY_CODE = "TCS"

# 2. Excel headers, by machines column (or helper name). Headers are matched trimmed and ignoring case.
HEADERS = {
    "az_no": "Az No",
    "asset_no": "TCS No",
    "main_category": "Main Category",
    "sub_category": "Sub Category",
    "sub_category_2": "Sub Category 2",
    "asset_type": "Type",
    "description": "Description",
    "manufacturer": "Manufacturer",
    "model": "Model Type",
    "capacity": "Capacity",
    "capacity_unit": "Capacity Unit",
    "company": "Company / Department",
    "ownership": "OwnerShip",
    "location": "Location",
    "project_name": "Project Name",
    "site": "Site",
    "mfg_year": "Mfg Year",
    "chassis": "Asset Chassis No.",
    "plate_no": "Plate No",
    "istimara_serial_no": "Istimara Serial No.",
    "insurance_expiry": "Insurance Expiry Date",
    "mvpi_expiry": "MVPI Exp. Date",
    "tuv_certificate_no": "TUV Certificate No",
    "tuv_certificate_expiry": "Certificate Exp. Date",
    "operational_card_no": "Operational Card No",
    "operational_card_expiry": "Operational card expiry date",
    "status_text": "Status",
    "mobilization_date": "Mobilization Date",
    "date_received": "Date Received",
    "fuel_type": "Fuel Type",
    "fuel_chip": "Fuel Chip",
    "gps_imei": "GPS IMEI Number",
    "gps_tracking": "GPS Tracking",
    "gps_date": "GPS Date",
    "rent_description": "Rent-Description",
}

# 3. Length limits, copied from the CHECK constraints on machines, customers and locations
TEXT_LIMITS = {
    "main_category": 100,
    "sub_category": 100,
    "sub_category_2": 100,
    "asset_type": 100,
    "description": 300,
    "manufacturer": 100,
    "model": 100,
    "capacity_unit": 100,
    "plate_no": 100,
    "project_name": 200,
    "istimara_serial_no": 100,
    "tuv_certificate_no": 100,
    "operational_card_no": 100,
    "fuel_type": 100,
    "fuel_chip": 100,
    "gps_imei": 100,
    "az_no": 50,
    "ownership_note": 200,
    "site": 200,
}
NAME_LIMIT = 200
DATE_FIELDS = [
    "insurance_expiry", "mvpi_expiry", "tuv_certificate_expiry", "operational_card_expiry",
    "mobilization_date", "date_received", "gps_date",
]

# 4. Value rules
NULL_TOKENS = {"", "NA", "N/A", "#N/A", "-", "0"}
SPACE_OR_CONTROL_RE = re.compile(r"[\s\x00-\x1f\x7f-\x9f]+")
ASSET_NO_RE = re.compile(r"^[A-Z0-9][A-Z0-9-]{0,29}$")
IDENTIFIER_RE = re.compile(r"^[A-Z0-9][A-Z0-9/.-]{0,39}$")
SERIAL_PREFIX_RE = re.compile(r"^S\s*[/.]\s*NO[^A-Z0-9]*")
MIN_YEAR = 1950
MAX_YEAR = datetime.date.today().year + 1
MAX_CAPACITY = Decimal("9999999999.99")  # numeric(12,2)

SITE_NO_TEXT = {"WORKSHOP REPAIRING", "SCRAP IN YARD"}
SISTER_VOTES = {"SISTER COMPANIES": True, "AZMEEL": True, "CUSTOMERS": False}
ALWAYS_SISTER = {"AZMEEL CONTRACTING COMPANY"}
STATUS_CODES = [
    "ready_in_yard", "in_yard_makeup", "under_repair", "awaiting_repair", "accident_damage",
    "broken_down_on_site", "on_hire_customer", "on_hire_sister", "sister_no_access",
    "azmeel_private", "in_tcs_use", "scrapped", "sold", "stolen", "missing",
]

# 5. Counts the current Total_Assets.xlsx must give (checked with Operations). --apply stops on a mismatch unless --force.
EXPECTED_COUNTS = {
    "total": 3767,
    "in_yard_makeup": 1971,
    "scrapped": 623,
    "under_repair": 381,
    "on_hire_sister": 259,
    "on_hire_customer": 212,
    "sister_no_access": 112,
    "ready_in_yard": 98,
    "azmeel_private": 71,
    "accident_damage": 21,
    "in_tcs_use": 19,
    "rejected": 0,
}

# 6. SQL for the machine upsert.
# A machine with checkins keeps its current status_id, and is_active follows that status.
# The rule is in SQL, so it also holds for a checkin made after the dry run.
MACHINE_COLUMNS = [
    "az_no", "main_category", "sub_category", "sub_category_2", "asset_type", "description",
    "manufacturer", "model", "capacity", "capacity_unit", "mfg_year", "chassis_no", "serial_no",
    "plate_no", "ownership_note", "project_name", "istimara_serial_no", "insurance_expiry",
    "mvpi_expiry", "tuv_certificate_no", "tuv_certificate_expiry", "operational_card_no",
    "operational_card_expiry", "mobilization_date", "date_received", "fuel_type", "fuel_chip",
    "gps_imei", "gps_installed", "gps_date", "status_id", "customer_id", "location_id", "site",
    "is_active",
]
HAS_CHECKINS_SQL = "exists (select 1 from public.status_checkins c where c.machine_id = m.id)"
KEEP_IF_CHECKINS = {
    "status_id": "m.status_id",
    "is_active": "(select s.is_active from public.status_values s where s.id = m.status_id)",
}
UPSERT_SQL = (
    "insert into public.machines as m (company_id, asset_no, "
    + ", ".join(MACHINE_COLUMNS)
    + ") values (%(company_id)s, %(asset_no)s, "
    + ", ".join(f"%({c})s" for c in MACHINE_COLUMNS)
    + ") on conflict (company_id, asset_no) do update set "
    + ", ".join(
        f"{c} = case when {HAS_CHECKINS_SQL} then {KEEP_IF_CHECKINS[c]} else excluded.{c} end"
        if c in KEEP_IF_CHECKINS
        else f"{c} = excluded.{c}"
        for c in MACHINE_COLUMNS
    )
)


class ImportStop(Exception):
    """A problem that stops the import before anything is written."""


class Report:
    """Collects the lines for data/import_report.csv."""

    def __init__(self):
        self.lines = []

    def add(self, asset_no, field, original, result, reason):
        self.lines.append({
            "asset_no": asset_no or "",
            "field": field,
            "original_value": "" if original is None else str(original),
            "result": "" if result is None else str(result),
            "reason": reason,
        })

    def reasons(self):
        return Counter(line["reason"] for line in self.lines)


# ---------- Cleaning helpers ----------

def clean_text(value):
    """Trim, collapse spaces and control characters to one space, turn null tokens into None."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime.datetime, datetime.date)):
        value = value.isoformat()
    text = SPACE_OR_CONTROL_RE.sub(" ", str(value)).strip()
    if text.upper() in NULL_TOKENS:
        return None
    return text


def upper_text(value):
    text = clean_text(value)
    return text.upper() if text else ""


def title_case(text):
    return " ".join(word.capitalize() for word in text.split(" "))


def clean_asset_no(value):
    text = clean_text(value)
    if text is None:
        return None
    return text.upper().replace(".", "-")


def to_decimal(value):
    """Return a finite Decimal, or None if the value is not a number."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = Decimal(str(value))
    else:
        try:
            number = Decimal(value)
        except InvalidOperation:
            return None
    return number if number.is_finite() else None


def parse_year(value):
    """Return (year or None, reason or None)."""
    text = clean_text(value)
    if text is None:
        return None, None
    number = to_decimal(value if isinstance(value, (int, float)) else text)
    if number is None or number != number.to_integral_value():
        return None, "invalid_mfg_year"
    year = int(number)
    if not MIN_YEAR <= year <= MAX_YEAR:
        return None, "mfg_year_out_of_range"
    return year, None


def parse_capacity(value):
    """Return (capacity or None, reason or None)."""
    text = clean_text(value)
    if text is None:
        return None, None
    number = to_decimal(value if isinstance(value, (int, float)) else text)
    if number is None:
        return None, "invalid_capacity"
    number = number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if number <= 0 or number > MAX_CAPACITY:
        return None, "capacity_out_of_range"
    return number, None


def parse_date(value):
    """Return (date or None, reason or None). Only real Excel dates are accepted."""
    if isinstance(value, datetime.datetime):
        return value.date(), None
    if isinstance(value, datetime.date):
        return value, None
    if clean_text(value) is None:
        return None, None
    return None, "invalid_date"


def split_chassis(value):
    """Return (field, value or None, reason or None). field is chassis_no or serial_no."""
    text = clean_text(value)
    if text is None:
        return "chassis_no", None, None
    text = text.upper()
    prefix = SERIAL_PREFIX_RE.match(text)
    if prefix:
        field, rest = "serial_no", text[prefix.end():]
    else:
        field, rest = "chassis_no", text
    rest = rest.replace(" ", "")
    if not rest:
        return field, None, "prefix_only"
    if not IDENTIFIER_RE.match(rest):
        return field, None, f"invalid_{field}"
    return field, rest, None


def status_for(asset_no, row):
    """Status code from the register, first rule that matches. None means the row is rejected."""
    rent = upper_text(row["rent_description"])
    site = upper_text(row["site"])
    status = upper_text(row["status_text"])
    if asset_no == "04-04-0028":
        return "in_tcs_use"
    if asset_no == "02-02-0025":
        return "scrapped"
    if rent == "SCRAP":
        return "scrapped"
    if site == "ACCIDENT":
        return "accident_damage"
    if rent in ("MAJOR BREAKDOWN", "FIXED WATER TANKS"):
        return "under_repair"
    if rent in ("GENERAL SERVICES", "TCS SITE USE", "INDUSTRIAL SAFETY"):
        return "in_tcs_use"
    if rent == "PRIVATE SECTOR":
        return "azmeel_private"
    if status == "ON JOB / NO ACCESS":
        return "sister_no_access"
    if rent == "CUSTOMERS":
        return "on_hire_customer"
    if rent in ("SISTER COMPANIES", "AZMEEL"):
        return "on_hire_sister"
    if rent == "YARD" and status == "AVAILABLE":
        return "ready_in_yard"
    if rent == "YARD":
        return "in_yard_makeup"
    return None


# ---------- Reading ----------

def open_sheet(path, sheet_name):
    if not path.exists():
        raise ImportStop(f"File not found: {path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ImportStop(f"Sheet {sheet_name!r} not found in {path.name}")
    return workbook[sheet_name]


def read_rows(path, sheet_name, headers):
    """Yield (excel_row_number, {name: value}) for each non-empty row. Columns are found by header."""
    sheet = open_sheet(path, sheet_name)
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None) or ()
    positions = {}
    for index, cell in enumerate(header_row):
        key = (clean_text(cell) or "").lower()
        positions.setdefault(key, index)
    missing = [h for h in headers.values() if h.lower() not in positions]
    if missing:
        raise ImportStop(f"{path.name}: header(s) not found: {', '.join(missing)}")
    columns = {name: positions[header.lower()] for name, header in headers.items()}
    for number, row in enumerate(rows, start=2):
        if all(clean_text(cell) is None for cell in row):
            continue
        yield number, {name: (row[i] if i < len(row) else None) for name, i in columns.items()}


def load_db_state(cur):
    cur.execute("select id, is_active from public.companies where short_code = %s", (COMPANY_CODE,))
    company = cur.fetchone()
    if company is None:
        raise ImportStop(f"Company {COMPANY_CODE} not found")
    if not company[1]:
        raise ImportStop(f"Company {COMPANY_CODE} is inactive")
    company_id = company[0]

    cur.execute("select code, id, is_active from public.status_values where company_id = %s", (company_id,))
    statuses = {code: (status_id, active) for code, status_id, active in cur.fetchall()}
    missing = [code for code in STATUS_CODES if code not in statuses]
    if missing:
        raise ImportStop(f"Status values missing for {COMPANY_CODE}: {', '.join(missing)}")

    cur.execute(
        "select id, name, is_sister_company, is_active from public.customers where company_id = %s",
        (company_id,),
    )
    customers = {
        name.lower(): {"id": cid, "name": name, "is_sister": sister, "is_active": active}
        for cid, name, sister, active in cur.fetchall()
    }

    cur.execute("select id, name, is_active from public.locations where company_id = %s", (company_id,))
    locations = {
        name.lower(): {"id": lid, "name": name, "is_active": active}
        for lid, name, active in cur.fetchall()
    }

    cur.execute(
        "select m.asset_no, s.code, "
        f"{HAS_CHECKINS_SQL} "
        "from public.machines m join public.status_values s on s.id = m.status_id "
        "where m.company_id = %s",
        (company_id,),
    )
    machines = {
        asset_no: {"status": code, "has_checkins": has_checkins}
        for asset_no, code, has_checkins in cur.fetchall()
    }

    return {
        "company_id": company_id,
        "status_ids": {code: value[0] for code, value in statuses.items()},
        "status_active": {code: value[1] for code, value in statuses.items()},
        "customers": customers,
        "locations": locations,
        "machines": machines,
    }


# ---------- Planning ----------

def limit_text(text, field, asset_no, report):
    if text is not None and len(text) > TEXT_LIMITS[field]:
        report.add(asset_no, field, text, None, "too_long")
        return None
    return text


def clean_name(value, field, asset_no, report):
    """Cleaned name for a customer or location, or None."""
    text = clean_text(value)
    if text is None:
        return None
    if len(text) > NAME_LIMIT:
        report.add(asset_no, field, text, None, "too_long")
        return None
    return text


def clean_machine(row, asset_no, status, report):
    machine = {"asset_no": asset_no}
    for field in ("main_category", "sub_category", "sub_category_2", "asset_type", "description",
                  "manufacturer", "model", "capacity_unit", "plate_no", "project_name",
                  "istimara_serial_no", "tuv_certificate_no", "operational_card_no", "fuel_type",
                  "fuel_chip", "gps_imei", "az_no"):
        machine[field] = limit_text(clean_text(row[field]), field, asset_no, report)

    ownership = clean_text(row["ownership"])
    machine["ownership_note"] = limit_text(title_case(ownership) if ownership else None,
                                           "ownership_note", asset_no, report)

    for field in DATE_FIELDS:
        machine[field], reason = parse_date(row[field])
        if reason:
            report.add(asset_no, field, row[field], None, reason)

    machine["gps_installed"] = True if upper_text(row["gps_tracking"]) == "INSTALLED" else None

    machine["capacity"], reason = parse_capacity(row["capacity"])
    if reason:
        report.add(asset_no, "capacity", row["capacity"], None, reason)

    machine["mfg_year"], reason = parse_year(row["mfg_year"])
    if reason:
        report.add(asset_no, "mfg_year", row["mfg_year"], None, reason)

    machine["chassis_no"] = machine["serial_no"] = None
    field, value, reason = split_chassis(row["chassis"])
    machine[field] = value
    if reason:
        report.add(asset_no, field, row["chassis"], None, reason)
    elif field == "serial_no":
        report.add(asset_no, "serial_no", row["chassis"], value, "moved_to_serial_no")

    # Customer: everything except TCS, name kept as typed
    company = clean_text(row["company"])
    customer = None
    if company is not None and company.upper() != COMPANY_CODE:
        customer = clean_name(company, "customer", asset_no, report)
    machine["customer_key"] = customer.lower() if customer else None
    machine["customer_name"] = customer

    location = clean_name(row["location"], "location", asset_no, report)
    location = title_case(location) if location else None
    machine["location_key"] = location.lower() if location else None
    machine["location_name"] = location

    # Site: kept as typed, except the two yard values that are not a site
    site = clean_text(row["site"])
    if site is not None and site.upper() in SITE_NO_TEXT:
        site = None
    machine["site"] = limit_text(site, "site", asset_no, report)

    machine["status"] = status
    machine["rent_description"] = upper_text(row["rent_description"])
    return machine


def null_duplicates(machines, field, report):
    """Set a value to None on every machine in the file that shares it."""
    counts = Counter(m[field] for m in machines if m[field] is not None)
    for machine in machines:
        value = machine[field]
        if value is not None and counts[value] > 1:
            report.add(machine["asset_no"], field, value, None, f"duplicate_{field}")
            machine[field] = None


def plan_customers(machines, state, report):
    """Customer names, sister flags, and what to create, rename or delete."""
    spellings, votes = {}, {}
    for machine in machines:
        key = machine["customer_key"]
        if key:
            spellings.setdefault(key, Counter())[machine["customer_name"]] += 1
            vote = SISTER_VOTES.get(machine["rent_description"])
            votes.setdefault(key, Counter())
            if vote is not None:
                votes[key][vote] += 1

    customers = {}
    for key, names in spellings.items():
        name = names.most_common(1)[0][0]  # most used spelling, first seen on a tie
        sister, customer = votes[key][True], votes[key][False]
        is_sister = name.upper() in ALWAYS_SISTER or sister > customer
        customers[key] = {"name": name, "is_sister": is_sister, "sister_votes": sister, "customer_votes": customer}

    new_customers, updates = {}, {}
    for key, c in customers.items():
        current = state["customers"].get(key)
        if current is None:
            new_customers[key] = c
            continue
        if current["name"] != c["name"] or current["is_sister"] != c["is_sister"]:
            updates[key] = c
            if current["name"] != c["name"]:
                report.add("", "customer.name", current["name"], c["name"], "customer_renamed")
            if current["is_sister"] != c["is_sister"]:
                report.add("", "customer.is_sister_company", f"{c['name']}: {current['is_sister']}",
                           c["is_sister"], "sister_flag_changed")
        if not current["is_active"]:
            report.add("", "customer", current["name"], current["name"], "inactive_customer")

    unused = sorted(c["name"] for key, c in state["customers"].items() if key not in customers)
    return customers, new_customers, updates, unused


def build_plan(state, report):
    stats = Counter()
    raw_rows = []
    for number, row in read_rows(REGISTER_FILE, REGISTER_SHEET, HEADERS):
        stats["rows_read"] += 1
        raw_rows.append((number, row, clean_asset_no(row["asset_no"])))

    # Reject blank, invalid and duplicate asset numbers, and rows that match no status rule
    asset_counts = Counter(a for _, _, a in raw_rows if a is not None)
    in_file = set()
    machines = []
    for number, row, asset_no in raw_rows:
        original = clean_text(row["asset_no"])
        if asset_no is None:
            report.add(f"(row {number})", "asset_no", row["asset_no"], "rejected", "blank_asset_no")
        elif not ASSET_NO_RE.match(asset_no):
            report.add(f"(row {number})", "asset_no", row["asset_no"], "rejected", "invalid_asset_no")
        else:
            in_file.add(asset_no)
            status = status_for(asset_no, row)
            if asset_counts[asset_no] > 1:
                report.add(asset_no, "asset_no", f"{original} (row {number})", "rejected", "duplicate_asset_no")
            elif status is None:
                shown = f"Rent-Description={clean_text(row['rent_description'])}; Status={clean_text(row['status_text'])}"
                report.add(asset_no, "status", shown, "rejected", "no_status_rule")
            else:
                if original != asset_no:
                    report.add(asset_no, "asset_no", original, asset_no, "asset_no_normalised")
                machines.append(clean_machine(row, asset_no, status, report))
                continue
        stats["rows_rejected"] += 1

    for field in ("chassis_no", "serial_no", "az_no"):
        null_duplicates(machines, field, report)

    # Machines in the database that are not in the file are deleted. A checkin blocks that.
    to_delete = sorted(a for a in state["machines"] if a not in in_file)
    blocked = [a for a in to_delete if state["machines"][a]["has_checkins"]]
    if blocked:
        raise ImportStop(
            f"{len(blocked)} machine(s) not in the file have status checkins and cannot be deleted: "
            + ", ".join(blocked)
        )

    # Machines with checkins keep their current status; is_active follows the status
    for machine in machines:
        current = state["machines"].get(machine["asset_no"])
        if current and current["has_checkins"]:
            if machine["status"] != current["status"]:
                report.add(machine["asset_no"], "status", machine["status"], current["status"], "has_checkins")
            stats["kept_for_checkins"] += 1
            machine["status"] = current["status"]
        machine["is_active"] = state["status_active"][machine["status"]]

    customers, new_customers, customer_updates, unused_customers = plan_customers(machines, state, report)

    new_locations = {}
    used_locations = set()
    for machine in machines:
        key, name = machine["location_key"], machine["location_name"]
        if key:
            used_locations.add(key)
            current = state["locations"].get(key)
            if current is None:
                new_locations.setdefault(key, name)
            elif not current["is_active"]:
                report.add(machine["asset_no"], "location", current["name"], current["name"], "inactive_location")
    unused_locations = sorted(loc["name"] for key, loc in state["locations"].items() if key not in used_locations)

    stats["to_insert"] = sum(1 for m in machines if m["asset_no"] not in state["machines"])
    stats["to_update"] = len(machines) - stats["to_insert"]
    return {
        "machines": machines,
        "to_delete": to_delete,
        "customers": customers,
        "new_customers": new_customers,
        "customer_updates": customer_updates,
        "unused_customers": unused_customers,
        "new_locations": new_locations,
        "unused_locations": unused_locations,
        "stats": stats,
    }


def count_mismatches(plan):
    """Return MISMATCH lines for every count that differs from EXPECTED_COUNTS."""
    status_counts = Counter(m["status"] for m in plan["machines"])
    actual = {
        "total": plan["stats"]["rows_read"],
        "rejected": plan["stats"]["rows_rejected"],
    }
    lines = []
    for name, expected in EXPECTED_COUNTS.items():
        got = actual[name] if name in actual else status_counts[name]
        if got != expected:
            lines.append(f"MISMATCH: {name} expected {expected}, got {got}")
    for code in sorted(set(status_counts) - set(EXPECTED_COUNTS)):
        lines.append(f"MISMATCH: {code} expected 0, got {status_counts[code]}")
    return lines


# ---------- Output ----------

def write_report(report):
    try:
        with open(REPORT_FILE, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(
                handle, fieldnames=["asset_no", "field", "original_value", "result", "reason"]
            )
            writer.writeheader()
            writer.writerows(report.lines)
    except PermissionError:
        raise ImportStop(f"Cannot write {REPORT_FILE}. Close it in Excel and try again.")


def print_summary(plan, report, mismatches):
    stats = plan["stats"]
    machines = plan["machines"]
    print("=== Register import summary ===")
    print(f"Rows read:                   {stats['rows_read']}")
    print(f"Rows rejected:               {stats['rows_rejected']}")
    print(f"Machines to insert:          {stats['to_insert']}")
    print(f"Machines to update:          {stats['to_update']}")
    print(f"Machines to delete:          {len(plan['to_delete'])} (in the database, not in the file)")
    print(f"Kept status (has checkins):  {stats['kept_for_checkins']}")
    print(f"Customers in the file:       {len(plan['customers'])}")
    for c in sorted(plan["customers"].values(), key=lambda c: c["name"].lower()):
        kind = "sister company" if c["is_sister"] else "customer"
        print(f"    {c['name']}  ({kind}; votes sister {c['sister_votes']}, customer {c['customer_votes']})")
    print(f"Customers to create:         {len(plan['new_customers'])}")
    print(f"Customers to rename or flag: {len(plan['customer_updates'])}")
    print(f"Customers to delete:         {len(plan['unused_customers'])}")
    for name in plan["unused_customers"]:
        print(f"    {name}")
    print(f"Locations to create:         {len(plan['new_locations'])}")
    for name in sorted(plan["new_locations"].values()):
        print(f"    {name}")
    print(f"Locations to delete:         {len(plan['unused_locations'])}")
    for name in plan["unused_locations"]:
        print(f"    {name}")
    print("Machines per status:")
    status_counts = Counter(m["status"] for m in machines)
    for code in STATUS_CODES:
        print(f"    {code:<20} {status_counts[code]}")
    print(f"Inactive machines:           {sum(1 for m in machines if not m['is_active'])}")
    print("Reported cleanups by reason:")
    for reason, count in sorted(report.reasons().items()):
        print(f"    {reason:<26} {count}")
    print("Expected counts check:")
    for line in mismatches or ["    Counts match."]:
        print(line)
    print(f"Report written to: {REPORT_FILE}")


# ---------- Writing ----------

def apply_plan(cur, plan, state):
    company_id = state["company_id"]

    # Delete machines that are not in the file first, so their chassis and serial numbers are free.
    if plan["to_delete"]:
        cur.execute(
            "delete from public.machines where company_id = %s and asset_no = any(%s)",
            (company_id, plan["to_delete"]),
        )

    location_ids = {key: loc["id"] for key, loc in state["locations"].items()}
    for key, name in plan["new_locations"].items():
        cur.execute(
            "insert into public.locations (company_id, name) values (%s, %s) returning id",
            (company_id, name),
        )
        location_ids[key] = cur.fetchone()[0]

    customer_ids = {key: c["id"] for key, c in state["customers"].items()}
    for key, c in plan["new_customers"].items():
        cur.execute(
            "insert into public.customers (company_id, name, is_sister_company) "
            "values (%s, %s, %s) returning id",
            (company_id, c["name"], c["is_sister"]),
        )
        customer_ids[key] = cur.fetchone()[0]

    for key, c in plan["customer_updates"].items():
        cur.execute(
            "update public.customers set name = %s, is_sister_company = %s where id = %s",
            (c["name"], c["is_sister"], customer_ids[key]),
        )

    # Clear the unique values first, so that two machines can swap a value
    # without a unique constraint error in the middle of the upsert.
    existing = [m["asset_no"] for m in plan["machines"] if m["asset_no"] in state["machines"]]
    if existing:
        cur.execute(
            "update public.machines set chassis_no = null, serial_no = null, az_no = null "
            "where company_id = %s and asset_no = any(%s)",
            (company_id, existing),
        )

    params = []
    for m in plan["machines"]:
        row = {c: m.get(c) for c in MACHINE_COLUMNS}
        row["company_id"] = company_id
        row["asset_no"] = m["asset_no"]
        row["status_id"] = state["status_ids"][m["status"]]
        row["customer_id"] = customer_ids[m["customer_key"]] if m["customer_key"] else None
        row["location_id"] = location_ids[m["location_key"]] if m["location_key"] else None
        params.append(row)
    cur.executemany(UPSERT_SQL, params)

    # Customers and locations that no machine uses any more
    for table, column in (("customers", "customer_id"), ("locations", "location_id")):
        cur.execute(
            f"delete from public.{table} t where t.company_id = %s "
            f"and not exists (select 1 from public.machines m where m.{column} = t.id)",
            (company_id,),
        )


# ---------- Main ----------

def safe_message(error, url):
    """Error text with the connection string removed."""
    return str(error).replace(url, "[DATABASE_URL]").strip()


def main():
    parser = argparse.ArgumentParser(description="Import the TCS machine register.")
    parser.add_argument("--apply", action="store_true", help="write to the database (default: dry run)")
    parser.add_argument("--force", action="store_true", help="with --apply: write even if the expected counts differ")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    load_dotenv(SCRIPT_DIR / ".env")
    url = os.environ.get("DATABASE_URL")
    if not url:
        print("Stopped: DATABASE_URL is not set in scripts/import_register/.env")
        return 1

    try:
        # prepare_threshold=None: no prepared statements, so the Supabase transaction pooler also works
        with psycopg.connect(url, connect_timeout=15, prepare_threshold=None) as conn:
            conn.isolation_level = psycopg.IsolationLevel.REPEATABLE_READ
            conn.read_only = not args.apply
            try:
                with conn.cursor() as cur:
                    state = load_db_state(cur)
                    report = Report()
                    plan = build_plan(state, report)
                    mismatches = count_mismatches(plan)
                    write_report(report)
                    print_summary(plan, report, mismatches)
                    if args.apply and mismatches and not args.force:
                        raise ImportStop("the counts above do not match. Nothing was written. "
                                         "Use --apply --force to write anyway.")
                    if args.apply:
                        apply_plan(cur, plan, state)
                if args.apply:
                    conn.commit()
                else:
                    conn.rollback()
            except BaseException:
                conn.rollback()
                raise
    except ImportStop as error:
        print(f"Stopped: {error}")
        return 1
    except psycopg.errors.ForeignKeyViolation as error:
        print("Stopped: a delete is blocked by a reference. Everything was rolled back.")
        print(f"    {safe_message(error, url)}")
        return 1
    except psycopg.Error as error:
        print(f"Database error ({type(error).__name__}), nothing was written: {safe_message(error, url)}")
        return 1

    if args.apply:
        print(f"APPLIED: {len(plan['machines'])} machines written, {len(plan['to_delete'])} deleted, "
              "in one transaction.")
    else:
        print("DRY RUN: nothing was written. Run again with --apply to write.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
